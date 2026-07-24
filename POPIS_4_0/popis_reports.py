from __future__ import annotations

from io import BytesIO
import math
import re
from typing import Any

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage


def _safe_sheet(name: str) -> str:
    return re.sub(r"[\\/*?:\[\]]", "_", str(name))[:31] or "Hoja"


def make_excel_report_with_images(tables: dict[str, pd.DataFrame], images: dict[str, bytes], context: dict[str, Any]) -> bytes:
    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        summary = pd.DataFrame([{"Campo": k, "Valor": v} for k, v in context.items() if not isinstance(v, (dict, list, tuple, pd.DataFrame))])
        summary.to_excel(writer, sheet_name="Resumen", index=False)
        for name, table in tables.items():
            if table is None:
                continue
            df = table if isinstance(table, pd.DataFrame) else pd.DataFrame(table)
            df.to_excel(writer, sheet_name=_safe_sheet(name), index=False)
        wb = writer.book
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row >= 1:
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="D9EAF7")
                    cell.alignment = Alignment(horizontal="center")
            for col in ws.columns:
                vals = [len(str(c.value)) if c.value is not None else 0 for c in col]
                ws.column_dimensions[col[0].column_letter].width = min(max(max(vals or [8]) + 2, 10), 45)
        if images:
            ws = wb.create_sheet("Graficas")
            row = 1
            for name, payload in images.items():
                ws.cell(row=row, column=1, value=name).font = Font(bold=True)
                img = XLImage(BytesIO(payload))
                max_width = 900
                if img.width > max_width:
                    ratio = max_width / img.width
                    img.width = max_width
                    img.height = int(img.height * ratio)
                ws.add_image(img, f"A{row+1}")
                row += max(24, int(img.height / 20) + 4)
    return bio.getvalue()


def _fmt(v, digits=1):
    try:
        if v is None or (isinstance(v, float) and math.isnan(v)):
            return "—"
        return f"{float(v):,.{digits}f}"
    except Exception:
        return str(v)


def make_word_report(context: dict[str, Any], tables: dict[str, pd.DataFrame], images: dict[str, bytes]) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    title = doc.add_heading(f"POPIS · Informe epidemiológico · {context.get('Año', '')} · SE{context.get('Semana', '')}", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Procesador Operativo de Patógenos e Indicadores Sanitarios")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Resumen ejecutivo", level=1)
    su = context.get("SUIVE acumulado")
    si = context.get("SINAVE acumulado")
    pos = context.get("SINAVE positivos")
    deaths = context.get("Defunciones registradas")
    state_inc = context.get("Incidencia estatal SINAVE")
    base = context.get("Base tasa", 100000)
    text = (
        f"Al corte de la semana epidemiológica {context.get('Semana')}, SUIVE registra {_fmt(su,0)} casos acumulados y "
        f"SINAVE {_fmt(si,0)} registros nominales. En SINAVE se identifican {_fmt(pos,0)} casos con patógeno y "
        f"{_fmt(deaths,0)} defunciones registradas. La incidencia estatal SINAVE es {_fmt(state_inc,2)} por {int(base):,} habitantes."
    )
    doc.add_paragraph(text)
    doc.add_paragraph("SUIVE y SINAVE son sistemas de vigilancia con universos distintos; sus conteos y tasas no deben sumarse ni interpretarse como equivalentes.")

    top = tables.get("Tasas municipal") or tables.get("Tasas municipales")
    if isinstance(top, pd.DataFrame) and not top.empty:
        inc_cols = [c for c in top.columns if str(c).startswith("Incidencia municipal")]
        ratio_cols = [c for c in top.columns if str(c).startswith("Razón incidencia municipal")]
        if inc_cols:
            t = top.sort_values(inc_cols[0], ascending=False).head(5)
            names = ", ".join(t["Municipio"].astype(str).tolist())
            doc.add_paragraph(f"Los municipios con mayor incidencia en el corte son: {names}. Estos territorios ameritan revisión conjunta con volumen de casos, tendencia semanal y concentración espacial.")
        if ratio_cols:
            high = top[pd.to_numeric(top[ratio_cols[0]], errors="coerce") > 1].sort_values(ratio_cols[0], ascending=False).head(5)
            if not high.empty:
                doc.add_paragraph("Los municipios con razón de incidencia superior a 1 presentan una incidencia mayor al promedio estatal y constituyen una señal descriptiva útil para priorizar vigilancia.")

    doc.add_heading("Tablas principales", level=1)
    preferred = ["Incidencia estatal", "Tasas municipal", "Patógenos", "SUIVE histórico", "Indicadores NuTraVE", "Indicadores laboratorio"]
    for name in preferred:
        df = tables.get(name)
        if not isinstance(df, pd.DataFrame) or df.empty:
            continue
        doc.add_heading(name, level=2)
        view = df.head(15)
        table = doc.add_table(rows=1, cols=len(view.columns))
        table.style = "Table Grid"
        for j, c in enumerate(view.columns):
            table.rows[0].cells[j].text = str(c)
        for _, row in view.iterrows():
            cells = table.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = "" if pd.isna(v) else str(v)

    if images:
        doc.add_heading("Gráficas", level=1)
        for name, payload in list(images.items())[:8]:
            doc.add_paragraph(name)
            doc.add_picture(BytesIO(payload), width=Inches(6.3))

    doc.add_heading("Fuentes y notas metodológicas", level=1)
    src = context.get("Fuentes", "Fuentes locales POPIS")
    doc.add_paragraph(str(src))
    doc.add_paragraph(
        "Incidencia municipal = casos SINAVE residentes del municipio / población proyectada del municipio × base seleccionada. "
        "Tasa con denominador estatal = casos del municipio / población proyectada total de Sonora × base. "
        "Razón municipio/estado = incidencia municipal / incidencia estatal SINAVE."
    )
    doc.add_paragraph("Las defunciones corresponden a registros con fecha de defunción en la base nominal y requieren validación de definición operacional para atribución definitiva a EDA.")

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
