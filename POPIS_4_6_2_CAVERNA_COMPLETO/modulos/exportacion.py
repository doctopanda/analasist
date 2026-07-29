from __future__ import annotations

import io
import zipfile
from datetime import datetime
from typing import Mapping

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt


def excel_bytes(tables: Mapping[str, pd.DataFrame]) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, df in tables.items():
            safe = str(name)[:31].replace("/", "_").replace("\\", "_").replace("?", "_").replace("*", "_")
            (df if isinstance(df, pd.DataFrame) else pd.DataFrame()).to_excel(writer, sheet_name=safe or "Hoja", index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = cell.font.copy(bold=True)
            for col in ws.columns:
                max_len = max((len(str(c.value or "")) for c in col[:1500]), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 11), 42)
    return output.getvalue()


def dataframe_to_doc_table(document: Document, df: pd.DataFrame, max_rows: int = 30) -> None:
    if df is None or df.empty:
        document.add_paragraph("Sin datos disponibles.")
        return
    view = df.head(max_rows).copy()
    table = document.add_table(rows=1, cols=len(view.columns))
    table.style = "Table Grid"
    for i, col in enumerate(view.columns):
        table.rows[0].cells[i].text = str(col)
    for _, row in view.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            if pd.isna(value):
                text = ""
            elif isinstance(value, float):
                text = f"{value:,.2f}"
            else:
                text = str(value)
            cells[i].text = text
    if len(df) > max_rows:
        document.add_paragraph(f"Se muestran {max_rows} de {len(df):,} filas. El Excel contiene el detalle completo.")


def word_bytes(title: str, tables: Mapping[str, pd.DataFrame], notes: list[str] | None = None) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    document.add_heading(title, level=0)
    p = document.add_paragraph("Procesador Operativo de Patógenos e Indicadores Sanitarios · POPIS 4.6.2")
    p.runs[0].font.size = Pt(10)
    document.add_paragraph(f"Generado: {datetime.now():%d/%m/%Y %H:%M}")
    for note in notes or []:
        document.add_paragraph(note, style="Intense Quote")
    for name, df in tables.items():
        document.add_heading(str(name), level=1)
        dataframe_to_doc_table(document, df)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def html_bytes(title: str, tables: Mapping[str, pd.DataFrame]) -> bytes:
    pieces = [
        "<!doctype html><html><head><meta charset='utf-8'>",
        f"<title>{title}</title>",
        "<style>body{font-family:Arial;background:#f4f6f9;color:#1f2937;max-width:1300px;margin:auto;padding:30px}h1,h2{color:#14213d}h2{border-left:5px solid #e76f76;padding-left:10px}table{border-collapse:collapse;width:100%;background:white;margin-bottom:25px}th{background:#14213d;color:white}th,td{border:1px solid #e5e7eb;padding:6px;font-size:12px}tr:nth-child(even){background:#fafafa}</style></head><body>",
        f"<h1>{title}</h1>",
    ]
    for name, df in tables.items():
        pieces.append(f"<h2>{name}</h2>")
        pieces.append(df.to_html(index=False, border=0) if isinstance(df, pd.DataFrame) and not df.empty else "<p>Sin datos.</p>")
    pieces.append("</body></html>")
    return "".join(pieces).encode("utf-8")


def _line_chart_png(df: pd.DataFrame, x: str, y: str, group: str | None, title: str) -> bytes:
    fig, ax = plt.subplots(figsize=(10, 5.2))
    if group and group in df.columns:
        for label, sub in df.groupby(group):
            ax.plot(sub[x], sub[y], marker="o", linewidth=1.5, markersize=2.5, label=str(label))
        ax.legend(frameon=False, ncol=3)
    else:
        ax.plot(df[x], df[y], marker="o", linewidth=1.5, markersize=2.5)
    ax.set_title(title, fontweight="bold")
    ax.set_xlabel(x)
    ax.set_ylabel(y)
    ax.grid(axis="y", alpha=.2)
    fig.tight_layout()
    output = io.BytesIO()
    fig.savefig(output, format="png", dpi=180, bbox_inches="tight")
    plt.close(fig)
    return output.getvalue()


def graphics_zip(weekly_sinave: pd.DataFrame | None = None, weekly_suive: pd.DataFrame | None = None,
                 pathogen_table: pd.DataFrame | None = None) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        if weekly_sinave is not None and not weekly_sinave.empty:
            zf.writestr("SINAVE_series_semanales.png", _line_chart_png(weekly_sinave, "Semana", "Casos", "Año", "SINAVE · casos por semana"))
        if weekly_suive is not None and not weekly_suive.empty:
            zf.writestr("SUIVE_series_semanales.png", _line_chart_png(weekly_suive, "Semana", "Casos", "Año", "SUIVE/SUAVE · casos por semana"))
        if pathogen_table is not None and not pathogen_table.empty:
            fig, ax = plt.subplots(figsize=(9, 5))
            view = pathogen_table.sort_values("Detecciones", ascending=True)
            ax.barh(view["Patógeno"], view["Detecciones"])
            ax.set_title("Patógenos identificados", fontweight="bold")
            ax.set_xlabel("Detecciones")
            fig.tight_layout()
            buf = io.BytesIO()
            fig.savefig(buf, format="png", dpi=180, bbox_inches="tight")
            plt.close(fig)
            zf.writestr("Patogenos.png", buf.getvalue())
        zf.writestr("LEEME.txt", "Gráficas generadas por POPIS 4.6.2 CAVERNA.\n")
    return output.getvalue()
